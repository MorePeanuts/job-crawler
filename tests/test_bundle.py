import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from job_crawler.bundle import bundle
from job_crawler.config import Settings
from job_crawler.models import Collection, Job, Scope, Source
from job_crawler.storage import write_json


def snapshot(path: Path, job: Job, now: datetime) -> Path:
    source = Source(source_id=job.source_id, company_id=job.company, company=job.company)
    write_json(
        path / "manifest.json",
        {
            "run_id": path.name,
            "as_of": now.isoformat(),
            "source_ids": [source.source_id],
            "bindings": {},
        },
    )
    item = path / "collections" / "source.json"
    write_json(
        item,
        {
            "source": source.model_dump(),
            "collection": Collection(
                jobs=[job],
                scopes=[
                    Scope(
                        name="jobs",
                        expected=1,
                        listed=1,
                        unique=1,
                        details=1,
                        terminal=True,
                    )
                ],
            ).model_dump(mode="json"),
        },
    )
    return item


async def test_merge_preserves_times_and_exports_review_rows(
    tmp_path: Path,
    job: Job,
    now: datetime,
) -> None:
    original_time = job.collected_at
    first = snapshot(tmp_path / "first", job, now)
    other = job.model_copy(update={"source_id": "OTHER", "deadline": "2000-01-01"})
    snapshot(tmp_path / "second", other, now)
    result = await bundle([first, tmp_path / "second"], tmp_path / "bundle", Settings())
    assert result["information_rows"] == 1 and result["review"] == 1
    assert len(result["sources"]) == 2
    assert result["sources"][1]["excluded"] == {"expired": 1}
    assert result["origins"][job.source_id]["as_of"] == now.isoformat()
    assert job.collected_at == original_time
    workbook = load_workbook(tmp_path / "bundle" / "岗位信息汇总.xlsx")
    assert workbook["岗位信息（待补全）"].max_row == 2
    assert workbook["岗位信息（待补全）"].max_column == 13
    assert workbook["企业采集情况"].max_row == 3
    assert workbook["企业采集情况"].cell(3, 7).value == 0
    workbook.close()
    chained = await bundle([tmp_path / "bundle"], tmp_path / "chained", Settings())
    assert chained["origins"][job.source_id]["as_of"] == now.isoformat()
    assert chained["information_rows"] == 1


async def test_merge_rejects_duplicate_source_and_incomplete_run(
    tmp_path: Path,
    job: Job,
    now: datetime,
) -> None:
    snapshot(tmp_path / "first", job, now)
    snapshot(tmp_path / "second", job, now)
    with pytest.raises(ValueError, match="来源重复"):
        await bundle([tmp_path / "first", tmp_path / "second"], tmp_path / "output", Settings())
    write_json(
        tmp_path / "first" / "manifest.json",
        {
            "run_id": "first",
            "as_of": now.isoformat(),
            "source_ids": [job.source_id, "MISSING"],
            "bindings": {},
        },
    )
    with pytest.raises(ValueError, match="缺少来源"):
        await bundle([tmp_path / "first"], tmp_path / "output", Settings())


async def test_daily_bundle_preserves_window_and_rejects_mixed_windows(
    tmp_path: Path,
    job: Job,
    now: datetime,
) -> None:
    old = job.model_copy(update={"published": "2026-01-01", "detail_mode": "listing_only"})
    daily = tmp_path / "daily"
    snapshot(daily, old, now)
    manifest = json.loads((daily / "manifest.json").read_text())
    manifest["window"] = "day"
    write_json(daily / "manifest.json", manifest)
    result = await bundle([daily], tmp_path / "result", Settings())
    assert result["window"] == "day" and result["as_of"] == now.isoformat()
    assert result["information_rows"] == 0
    assert result["sources"][0]["excluded"] == {"outside_window": 1}
    snapshot(tmp_path / "all", job.model_copy(update={"source_id": "OTHER"}), now)
    with pytest.raises(ValueError, match="不同时间窗口"):
        await bundle([daily, tmp_path / "all"], tmp_path / "mixed", Settings())


async def test_bundle_preserves_excel_incompatible_text_in_sidecar(
    tmp_path: Path, job: Job, now: datetime
) -> None:
    original = "岗位职责\x1e完整正文"
    snapshot(tmp_path / "source", job.model_copy(update={"duties": original}), now)
    result = await bundle([tmp_path / "source"], tmp_path / "result", Settings())
    assert result["information_overflow_fields"] == 1
    overflow = json.loads((tmp_path / "result" / "excel-overflow.json").read_text())
    assert overflow[0]["value"] == original
    workbook = load_workbook(tmp_path / "result" / "岗位信息汇总.xlsx")
    assert "excel-overflow.json" in workbook["岗位信息（待补全）"].cell(2, 9).value
    workbook.close()
