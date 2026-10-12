from pathlib import Path

import pytest
from openpyxl import load_workbook

from job_crawler.config import MappingConfig
from job_crawler.exporter import (
    HEADERS,
    enums,
    export_workbook,
    import_errors,
    normalize,
    validate_mapping,
)
from job_crawler.models import Job


def test_template_contract_and_formula_safety(tmp_path: Path, job: Job) -> None:
    job.title = '=HYPERLINK("https://example.test","bad")'
    path = tmp_path / "jobs.xlsx"
    export_workbook(path, [job])
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["岗位导入", "枚举值", "填写说明"]
    sheet = workbook["岗位导入"]
    assert [c.value for c in sheet[1]] == HEADERS
    assert sheet.max_row == 2
    assert sheet["B2"].value == job.title
    assert sheet["B2"].data_type == "s"
    assert sheet["A2"].value == "测试公司"
    assert len(sheet.data_validations.dataValidation) == 4
    assert workbook["枚举值"].max_row == 114
    workbook.close()


def test_empty_output_removes_examples(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    export_workbook(path, [])
    workbook = load_workbook(path)
    assert workbook["岗位导入"].max_row == 1
    assert workbook["岗位导入"]["B2"].value is None
    workbook.close()


@pytest.mark.parametrize("information", [False, True])
@pytest.mark.parametrize(
    ("deadline", "expected"),
    [
        ("2026-10-15", "2026-10-15"),
        ("2026-10-12T00:00:00", "2026-10-12"),
        ("2026-10-11T17:00:00Z", "2026-10-12"),
        ("2026年10月2日", "2026-10-02"),
        ("2026/10/2", "2026-10-02"),
        ("1970-01-01", "1970-01-01"),
        ("0000000000", "1970-01-01"),
        ("", ""),
        ("尚未公布", ""),
        ('=HYPERLINK("https://example.test","x")', ""),
    ],
)
def test_optional_deadline_is_last_column_and_preserves_source(
    tmp_path: Path,
    job: Job,
    information: bool,
    deadline: str,
    expected: str,
    mapping: MappingConfig,
) -> None:
    job.deadline = deadline
    path = tmp_path / "deadline.xlsx"
    export_workbook(path, [job], information=information)
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    assert sheet.max_column == 13
    assert sheet["K1"].value == "申请链接"
    assert sheet["L1"].value == "截止时间"
    assert sheet["M1"].value == "发布时间"
    assert (sheet["L2"].value or "") == expected
    assert job.deadline == deadline
    assert sheet["L2"].data_type != "f"
    assert sheet["L2"].number_format == "@"
    if information:
        assert sheet.auto_filter.ref == "A1:M2"
    if not deadline:
        assert import_errors(job, mapping, enums()) == []
    workbook.close()


@pytest.mark.parametrize("information", [False, True])
@pytest.mark.parametrize(
    ("published", "expected"),
    [
        ("2026-10-10", "2026-10-10"),
        ("2026-10-11T17:00:00Z", "2026-10-12"),
        ("2026年10月2日", "2026-10-02"),
        ("1791763200", "2026-10-12"),
        ("", ""),
        ("未公布", ""),
        ('=HYPERLINK("https://example.test","x")', ""),
    ],
)
def test_publication_is_last_column_and_preserves_source(
    tmp_path: Path, job: Job, information: bool, published: str, expected: str
) -> None:
    job.published = published
    path = tmp_path / "publication.xlsx"
    export_workbook(path, [job], information=information)
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    assert sheet.max_column == 13
    assert sheet["M1"].value == "发布时间"
    assert (sheet["M2"].value or "") == expected
    assert sheet["M2"].data_type != "f"
    assert sheet["M2"].number_format == "@"
    assert job.published == published
    workbook.close()


def test_category_tree_never_inferred_from_enum_rows(job: Job) -> None:
    assert "invalid:category_path" in import_errors(job, MappingConfig(), enums())
    with pytest.raises(ValueError, match="分类映射"):
        validate_mapping(MappingConfig(categories={"工程": ["互联网/AI", "软件研发"]}))


def test_exact_mapping_preserves_source(job: Job, mapping: MappingConfig) -> None:
    job.education = "本科及以上"
    job.location = "深圳"
    job.category_raw = "开发"
    mapping.education = {"本科及以上": "本科"}
    mapping.locations = {"深圳": "广东省/深圳市"}
    mapping.categories = {"开发": ["互联网/AI", "软件研发"]}
    validate_mapping(mapping)
    result = normalize(job, mapping)
    assert import_errors(result, mapping, enums()) == []
    assert job.education == "本科及以上"


def test_no_silent_truncation_or_unknown_enum(
    tmp_path: Path, job: Job, mapping: MappingConfig
) -> None:
    job.duties = "长" * 32768
    job.education = "不限"
    assert "invalid:education" in import_errors(job, mapping, enums())
    assert "invalid:excel_text:duties" in import_errors(job, mapping, enums())
    with pytest.raises(ValueError, match="无损"):
        export_workbook(tmp_path / "bad.xlsx", [job])
