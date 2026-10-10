"""Exercise one-command delivery with real adapters/exporters and simulated I/O."""

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx2 as httpx
import pytest
from openai import APITimeoutError, AsyncOpenAI
from openpyxl import load_workbook
from test_incremental import DailyReplay
from test_runner import complete

from job_crawler import runner
from job_crawler.classification import Classifier, job_content
from job_crawler.cli import main
from job_crawler.config import CrawlConfig, LLMConfig, MappingConfig, Settings, load_settings
from job_crawler.delivery import configured_taxonomy, enrich_delivery
from job_crawler.llm import Enricher
from job_crawler.models import Issue, Job, Source
from job_crawler.storage import Store, write_json


def model_stub(monkeypatch: pytest.MonkeyPatch, *, invalid: bool = False) -> list[str]:
    calls: list[str] = []

    def respond(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        system = body["messages"][0]["content"]
        if "岗位职能分类员" in system:
            calls.append("category")
            prompt = body["messages"][1]["content"]
            items = json.loads(prompt[prompt.index("[") :])
            answer = {
                "items": [
                    {
                        "id": item["id"],
                        "primary": "互联网/AI",
                        "secondary": "软件研发",
                        "quote": "not in source" if invalid else item["title"],
                        "reason": "软件研发职责",
                    }
                    for item in items
                ]
            }
        else:
            calls.append("education")
            answer = {"education": {"value": "本科", "quote": "学士学位"}}
        return httpx.Response(
            200,
            json={
                "id": "mock",
                "created": 0,
                "model": "test",
                "object": "chat.completion",
                "choices": [
                    {
                        "index": 0,
                        "finish_reason": "stop",
                        "message": {
                            "role": "assistant",
                            "content": json.dumps(answer),
                        },
                    }
                ],
            },
        )

    def client(**kwargs: Any) -> AsyncOpenAI:
        return AsyncOpenAI(
            **kwargs, http_client=httpx.AsyncClient(transport=httpx.MockTransport(respond))
        )

    monkeypatch.setenv("DEEPSEEK_API_KEY", "offline-test-only")
    monkeypatch.setattr("job_crawler.classification.AsyncOpenAI", client)
    monkeypatch.setattr("job_crawler.llm.AsyncOpenAI", client)
    return calls


def configure(tmp_path: Path) -> Path:
    text = Path("config.toml").read_text()
    text = text.replace(
        'cache_dir = ".runtime/classification"', f'cache_dir = "{tmp_path}/categories"'
    )
    text = text.replace(
        'education_cache_dir = ".runtime/education"',
        f'education_cache_dir = "{tmp_path}/education"',
    )
    path = tmp_path / "config.toml"
    path.write_text(text)
    return path


def browser_stub(monkeypatch: pytest.MonkeyPatch) -> list[DailyReplay]:
    sessions: list[DailyReplay] = []

    class Browser(DailyReplay):
        def __init__(
            self,
            source_id: str,
            run_id: str,
            output: Path,
            store: Store,
            config: CrawlConfig,
            resume: bool = False,
        ):
            output.mkdir(parents=True)
            super().__init__(output / "listing.json", count=1, dates={"0": "2026-10-10"})
            self.navigation_warnings: list[Issue] = []
            sessions.append(self)

        async def __aenter__(self) -> "Browser":
            return self

        async def __aexit__(self, *args: object) -> None:
            pass

        async def request(
            self, url: str, *, payload: dict[str, Any] | None = None, detail: bool = False
        ) -> dict[str, Any]:
            data = await super().request(url, payload=payload, detail=detail)
            if detail:
                data["Data"].update(
                    Require="须取得学士学位，负责软件研发。", LocNames="广东省/深圳市"
                )
            return data

    monkeypatch.setattr(runner, "BrowserSession", Browser)
    return sessions


@pytest.mark.parametrize("window", ["all", "day"])
def test_one_cli_command_delivers_and_next_run_uses_both_caches(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, window: str
) -> None:
    config = configure(tmp_path)
    browser_stub(monkeypatch)
    calls = model_stub(monkeypatch)
    args = [
        "run",
        "--source",
        "COM0056",
        "--config",
        str(config),
        "--window",
        window,
        "--as-of",
        "2026-10-10T23:59:59+08:00",
        "--state",
        str(tmp_path / "state.sqlite"),
    ]
    for name in ("first", "second"):
        output = tmp_path / name
        assert main([*args, "--output", str(output)]) == 0
        workbook = load_workbook(output / "jobs.xlsx")
        sheet = workbook.worksheets[0]
        assert sheet.max_column == 12 and sheet.max_row == 2
        assert [sheet.cell(2, col).value for col in (4, 5, 6)] == ["互联网/AI", "软件研发", "本科"]
        workbook.close()
        report = json.loads((output / "report.json").read_text())
        assert report["status"] == "complete" and report["review"] == 0
        assert report["enrichment"]["calls"] == (2 if name == "first" else 0)
        assert report["enrichment"]["cache_hits"] == (0 if name == "first" else 2)
        raw = json.loads((output / "jobs.json").read_text())[0]
        delivered = json.loads((output / "delivered.json").read_text())[0]
        assert raw["category_primary"] == raw["education"] == ""
        assert delivered["category_secondary"] == "软件研发"
    assert calls == ["category", "education"]


def test_missing_key_still_exports_and_export_retries_without_browser(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config = configure(tmp_path)
    sessions = browser_stub(monkeypatch)
    original = tmp_path / "original"
    assert (
        main(
            [
                "run",
                "--source",
                "COM0056",
                "--config",
                str(config),
                "--output",
                str(original),
                "--state",
                str(tmp_path / "state.sqlite"),
            ]
        )
        == 2
    )
    report = json.loads((original / "report.json").read_text())
    assert report["review"] == 1 and report["enrichment"]["calls"] == 0
    assert "DEEPSEEK_API_KEY" in (original / "review.json").read_text()
    assert (original / "jobs.json").exists() and (original / "岗位信息汇总.xlsx").exists()
    calls = model_stub(monkeypatch)
    result = tmp_path / "retried"
    assert main(["export", str(original), "--output", str(result), "--config", str(config)]) == 0
    assert len(sessions) == 1 and calls == ["category", "education"]
    saved = json.loads((result / "report.json").read_text())
    assert saved["as_of"] == report["as_of"] and saved["exported"] == 1


@pytest.mark.parametrize("invalid", [False, True])
async def test_global_budget_includes_classification_and_education(
    tmp_path: Path, job: Job, now: datetime, monkeypatch: pytest.MonkeyPatch, invalid: bool
) -> None:
    settings = load_settings(configure(tmp_path))
    settings.llm.max_calls = 1
    calls = model_stub(monkeypatch, invalid=invalid)
    job.category_primary = job.category_secondary = job.education = ""
    job.requirements = "须取得学士学位"
    errors = await enrich_delivery([job], settings, tmp_path / "result", now)
    assert calls == ["category"]
    assert "预算" in str(errors)
    assert not list((tmp_path / "education").glob("*.json"))
    if invalid:
        assert job.category_secondary == ""
        assert not list((tmp_path / "categories").glob("*.json"))
    else:
        assert job.category_secondary == "软件研发"


async def test_selection_and_explicit_values_before_llm(
    tmp_path: Path, job: Job, now: datetime, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = load_settings(configure(tmp_path))
    calls = model_stub(monkeypatch)
    unknown = job.model_copy(
        update={"source_key": "unknown", "published": "bad date", "category_secondary": ""}
    )
    old = job.model_copy(
        update={"source_key": "old", "published": "2020-01-01", "category_secondary": ""}
    )
    source = Source(source_id=job.source_id, company_id="TEST", company=job.company)
    valid, review, reports = await runner.prepare_export(
        {source.source_id: complete([job, unknown, old])},
        [source],
        {},
        settings,
        "day",
        now,
        tmp_path / "result",
    )
    assert not calls and len(valid) == len(review) == 1
    assert valid[0].category_secondary == "软件研发"
    assert reports[0].excluded == {"outside_window": 1}


@pytest.mark.parametrize("change", ["content", "model", "policy", "corrupt"])
async def test_education_cache_validated_and_invalidated(
    tmp_path: Path, job: Job, monkeypatch: pytest.MonkeyPatch, change: str
) -> None:
    calls = model_stub(monkeypatch)
    job.requirements = "须取得学士学位"
    job.education = ""
    config = LLMConfig()
    mapping = MappingConfig()
    worker = Enricher(
        config, mapping, tmp_path / "audit", education_only=True, cache=tmp_path / "cache"
    )
    try:
        result = await worker.enrich(job)
        assert result.education == "本科" and result.category_secondary == job.category_secondary
        config.max_calls = 0
        monkeypatch.delenv("DEEPSEEK_API_KEY")
        assert (await worker.enrich(job)).education == "本科"
        assert worker.cache_hits == 1 and calls == ["education"]
        if change == "content":
            job.requirements += "；计算机专业"
        elif change == "model":
            config.model = "another-model"
        elif change == "policy":
            mapping.education_policy = "explicit_only"
        else:
            saved = next((tmp_path / "cache").glob("*.json"))
            data = json.loads(saved.read_text())
            data["decision"]["education"]["quote"] = "虚构引证"
            saved.write_text(json.dumps(data))
        with pytest.raises(ValueError, match="预算"):
            await worker.enrich(job)
    finally:
        await worker.close()


async def test_disabled_pipeline_makes_no_model_requests(
    tmp_path: Path, job: Job, now: datetime, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = Settings(llm=LLMConfig(enabled=False))
    calls = model_stub(monkeypatch)
    assert not await enrich_delivery([job], settings, tmp_path, now)
    assert not calls


async def test_cached_refinement_does_not_retry_uncached_base_decision(
    tmp_path: Path, job: Job, now: datetime, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = load_settings(configure(tmp_path))
    calls = model_stub(monkeypatch)
    job.category_primary = job.category_secondary = ""
    reviewer = Classifier(
        settings,
        configured_taxonomy(settings),
        Path(settings.llm.cache_dir),
        tmp_path / "audit",
        max_calls=0,
        batch_size=1,
        concurrency=1,
        compact=True,
    )
    content = job_content(job)
    key = reviewer.content_id(content)
    write_json(
        reviewer.cache / f"{key}.json",
        {
            "identity": reviewer.identity,
            "input": content,
            "decision": {
                "id": key,
                "primary": "互联网/AI",
                "secondary": "软件研发",
                "quote": job.title,
                "reason": "职责明确",
            },
        },
    )
    assert not await enrich_delivery([job], settings, tmp_path / "result", now)
    assert not calls and job.category_secondary == "软件研发"
    report = json.loads((tmp_path / "result" / "enrichment-report.json").read_text())
    assert report["cache_hits"] == 1 and report["calls"] == 0


async def test_model_timeout_preserves_other_valid_jobs_and_exports_partial(
    tmp_path: Path, job: Job, now: datetime, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = load_settings(configure(tmp_path))
    settings.llm.max_calls = 1
    needs_education = job.model_copy(
        update={
            "source_key": "needs-education",
            "education": "",
            "requirements": "须取得学士学位",
        }
    )
    collections = {job.source_id: complete([job, needs_education])}
    source = Source(source_id=job.source_id, company_id="TEST", company=job.company)

    def respond(request: httpx.Request) -> httpx.Response:
        raise APITimeoutError(request=request)

    def client(**kwargs: Any) -> AsyncOpenAI:
        return AsyncOpenAI(
            **kwargs, http_client=httpx.AsyncClient(transport=httpx.MockTransport(respond))
        )

    monkeypatch.setenv("DEEPSEEK_API_KEY", "do-not-persist-secret")
    monkeypatch.setattr("job_crawler.llm.AsyncOpenAI", client)
    valid, review, reports = await runner.prepare_export(
        collections, [source], {}, settings, "all", now, tmp_path / "result"
    )
    report = runner.save_results(tmp_path / "result", {}, collections, valid, review, reports)
    assert report["status"] == "partial" and report["exported"] == report["review"] == 1
    assert report["enrichment"]["calls"] == 1
    assert "APITimeoutError" in str(review[0]["reasons"])
    assert "do-not-persist-secret" not in (tmp_path / "result" / "report.json").read_text()
    assert (tmp_path / "result" / "jobs.xlsx").exists()
