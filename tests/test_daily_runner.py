import json
from datetime import datetime, timedelta, tzinfo
from pathlib import Path
from typing import Self

import pytest
from openpyxl import load_workbook
from test_incremental import DailyReplay
from test_zhiye import source

from job_crawler import runner
from job_crawler.adapters.zhiye import Xiangyu
from job_crawler.config import CrawlConfig, Settings
from job_crawler.models import Binding, Issue
from job_crawler.storage import Store


async def test_run_persists_baseline_and_exports_only_undated_new_jobs(
    tmp_path: Path,
    now: datetime,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    clock = now
    sessions: list[DailyReplay] = []
    count = 1

    class Clock(datetime):
        @classmethod
        def now(cls, tz: tzinfo | None = None) -> Self:
            return cls.fromisoformat(clock.isoformat())

    class Browser(DailyReplay):
        def __init__(
            self,
            source_id: str,
            run_id: str,
            output: Path,
            store: Store,
            config: CrawlConfig,
            resume: bool = False,
        ):
            output.mkdir(parents=True)
            super().__init__(output / "list.json", count)
            self.navigation_warnings: list[Issue] = []
            sessions.append(self)

        async def __aenter__(self) -> "Browser":
            return self

        async def __aexit__(self, *args: object) -> None:
            pass

    monkeypatch.setattr(runner, "datetime", Clock)
    monkeypatch.setattr(runner, "BrowserSession", Browser)
    sources = [source()]
    bindings = {sources[0].source_id: Binding(adapter=Xiangyu.name)}
    settings = Settings()
    state = tmp_path / "state.sqlite"
    baseline = await runner.run(
        sources, bindings, settings, output=tmp_path / "base", state=state, window="all"
    )
    assert baseline["sources"][0]["coverage"] == "complete"
    clock += timedelta(days=1)
    count = 2
    output = tmp_path / "daily"
    report = await runner.run(sources, bindings, settings, output=output, state=state, window="day")
    assert sessions[1].details == ["1"]
    assert report["sources"][0]["scopes"][0]["reused_details"] == 1
    assert report["sources"][0]["selection_basis"] == {"first_discovered_without_publication": 1}
    rows = json.loads((output / "review.json").read_text())
    assert len(rows) == 1 and rows[0]["job"]["source_key"] == "1"
    assert rows[0]["job"]["published"] == ""
    assert "missing_publication" not in rows[0]["reasons"]
    workbook = load_workbook(output / "岗位信息汇总.xlsx")
    assert workbook.worksheets[0].max_row == 2
    assert workbook.worksheets[0]["B2"].value == "工程师1"
    workbook.close()
    reexported = await runner.reexport(output, tmp_path / "reexport", settings)
    assert reexported["as_of"] == report["as_of"] and reexported["window"] == "day"
    assert reexported["review"] == 1
