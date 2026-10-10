"""Compare a historical workbook with its field-enriched offline re-export."""

import argparse
import json
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.escape import unescape

from job_crawler.models import Job
from job_crawler.storage import write_json


def read_json(path: Path) -> object:
    return json.loads(path.read_text())


def validations(path: Path) -> list[tuple[str, str, str]]:
    """Excel may store cross-sheet lists as x14 extensions instead of main XML."""
    result: list[tuple[str, str, str]] = []
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
                continue
            for element in ET.fromstring(archive.read(name)).iter():
                if element.tag.rsplit("}", 1)[-1] != "dataValidation":
                    continue
                fields = {
                    child.tag.rsplit("}", 1)[-1]: "".join(child.itertext()) for child in element
                }
                result.append(
                    (
                        name,
                        element.get("sqref", fields.get("sqref", "")),
                        fields.get("formula1", "").replace("'枚举值'!", "枚举值!"),
                    )
                )
    return sorted(result)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("before", type=Path)
    parser.add_argument("after", type=Path)
    args = parser.parse_args()
    before = load_workbook(args.before / "岗位信息汇总.xlsx")
    after = load_workbook(args.after / "岗位信息汇总.xlsx")
    changes: Counter[str] = Counter()
    assert before.sheetnames == after.sheetnames
    assert validations(args.before / "岗位信息汇总.xlsx") == validations(
        args.after / "岗位信息汇总.xlsx"
    )
    for old, new in zip(before, after, strict=True):
        assert (old.max_row, old.max_column) == (new.max_row, new.max_column)
        for old_row, new_row in zip(old, new, strict=True):
            for a, b in zip(old_row, new_row, strict=True):
                old_value = unescape(a.value) if isinstance(a.value, str) else a.value
                new_value = unescape(b.value) if isinstance(b.value, str) else b.value
                if old_value == new_value:
                    continue
                allowed = (old == before.worksheets[0] and a.row > 1 and a.column in (6, 8)) or (
                    old.title == "企业采集情况" and a.row > 1 and a.column == 8
                )
                assert allowed, (old.title, a.coordinate, a.value, b.value)
                changes[f"{old.title}:{a.column_letter}"] += 1
    before.close()
    after.close()
    assert read_json(args.before / "jobs.json") == read_json(args.after / "jobs.json")
    old_manifest = json.loads((args.before / "manifest.json").read_text())
    new_manifest = json.loads((args.after / "manifest.json").read_text())
    assert old_manifest["as_of"] == new_manifest["as_of"]
    for source, origin in old_manifest["origins"].items():
        assert new_manifest["origins"][source]["as_of"] == origin["as_of"]
    originals = {
        j.key: j
        for j in (
            Job.model_validate(raw) for raw in json.loads((args.after / "jobs.json").read_text())
        )
    }
    audit = json.loads((args.after / "field-audit.json").read_text())
    verified_quotes = 0
    for record in audit["records"]:
        original = originals[record["key"]]
        for field in record["fields"].values():
            if field["value"] and field["value"] != field["original"]:
                prefix, quote = field["provenance"].split(";quote:", 1)
                source_field = prefix.split(";field:")[1]
                assert quote and quote in getattr(original, source_field)
                verified_quotes += 1
    raw_audit = json.loads((args.after / "saved-evidence-field-audit.json").read_text())
    assert raw_audit["statuses"] == {"matched": len(audit["records"])}
    result = {
        "passed": True,
        "changed_cells": dict(changes),
        "verified_change_quotes": verified_quotes,
        "source_jobs_unchanged": len(originals),
        "original_times_preserved": True,
        "summary": audit["summary"],
    }
    write_json(args.after / "field-verification.json", result)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
