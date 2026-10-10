"""Check a classified workbook against the original delivered snapshot, offline."""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from job_crawler.config import load_settings
from job_crawler.exporter import FIELDS, HEADERS
from job_crawler.filtering import select
from job_crawler.models import Job
from job_crawler.storage import write_json


def read_information(root: Path) -> list[Job]:
    report = json.loads((root / "review.json").read_text())
    review = {
        item["job"]["source_id"] + ":" + item["job"]["source_key"]: Job.model_validate(item["job"])
        for item in report
    }
    # Information rows are normalized review records plus import-valid records.
    from datetime import datetime

    from job_crawler.exporter import normalize

    settings = load_settings(Path("config.toml"))
    manifest = json.loads((root / "manifest.json").read_text())
    for item in json.loads((root / "jobs.json").read_text()):
        job = Job.model_validate(item)
        if (
            job.key not in review
            and select(job, "all", datetime.fromisoformat(manifest["as_of"])).kind != "exclude"
        ):
            review[job.key] = normalize(job, settings.mapping)
    return sorted(review.values(), key=lambda j: (j.company, j.title, j.key))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("original", type=Path)
    parser.add_argument(
        "classified", type=Path, help="classify output directory (parent of result)"
    )
    args = parser.parse_args()
    result = args.classified / "result"
    before = read_information(args.original)
    after = read_information(result)
    assert [j.key for j in before] == [j.key for j in after], "岗位集合变化"
    tree = load_settings(Path("config.toml")).mapping.category_paths
    for old, new in zip(before, after, strict=True):
        for field in FIELDS:
            if field not in {"category_primary", "category_secondary"}:
                assert getattr(old, field) == getattr(new, field), (new.key, field)
        if new.category_primary or new.category_secondary:
            assert new.category_secondary in tree.get(new.category_primary, [])
    workbook = load_workbook(result / "岗位信息汇总.xlsx")
    sheet = workbook.worksheets[0]
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.max_row == len(after) + 1
    for job, row in zip(after, sheet.iter_rows(min_row=2), strict=True):
        for field, cell in zip(FIELDS, row, strict=True):
            assert (cell.value or "") == getattr(job, field), (job.key, field)
            assert cell.data_type != "f", "文本被写为公式"
    workbook.close()
    old_report = json.loads((args.original / "report.json").read_text())
    new_report = json.loads((result / "report.json").read_text())
    old_sources = {row["source_id"]: row for row in old_report["sources"]}
    new_sources = {row["source_id"]: row for row in new_report["sources"]}
    assert old_sources.keys() == new_sources.keys()
    for source_id, old in old_sources.items():
        new = new_sources[source_id]
        for field in ("source_id", "coverage", "collected", "excluded"):
            assert old[field] == new[field], field
    for source_id, origin in old_report["origins"].items():
        assert origin["as_of"] == new_report["origins"][source_id]["as_of"]
    summary = json.loads((args.classified / "classification-report.json").read_text())
    assert summary["classified"] == sum(bool(j.category_secondary) for j in after)
    audit = {item["job"]: item for item in summary["items"]}
    for job in after:
        item = audit[job.key]
        if job.category_secondary:
            assert (item["decision"]["primary"], item["decision"]["secondary"]) == (
                job.category_primary,
                job.category_secondary,
            )
    validation = {
        "rows": len(after),
        "sources": len(new_report["sources"]),
        "classified": summary["classified"],
        "unresolved": summary["unresolved"],
        "unchanged_other_fields": True,
        "source_times_preserved": True,
        "workbook_matches_json": True,
    }
    write_json(args.classified / "verification.json", validation)
    print(json.dumps(validation, ensure_ascii=False))


if __name__ == "__main__":
    main()
