"""Read-only batch delivery checks. Never fetch sites or call a model."""

import argparse
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import cast

from openpyxl import load_workbook

from job_crawler.filtering import Window, select
from job_crawler.models import Collection, Job, Source, SourceReport
from job_crawler.storage import write_json


def audit(root: Path, expected_companies: int | None = None) -> dict[str, object]:
    manifest = json.loads((root / "manifest.json").read_text())
    report = json.loads((root / "report.json").read_text())
    reports = [SourceReport.model_validate(row) for row in report["sources"]]
    problems: list[str] = []
    source_ids = [row.source_id for row in reports]
    if len(source_ids) != len(set(source_ids)) or set(source_ids) != set(manifest["source_ids"]):
        problems.append("报告来源重复或与 manifest 不一致")
    sources: dict[str, Source] = {}
    collections: dict[str, Collection] = {}
    for path in sorted((root / "collections").glob("*.json")):
        raw = json.loads(path.read_text())
        source = Source.model_validate(raw["source"])
        if source.source_id in sources:
            problems.append(f"快照来源重复: {source.source_id}")
        sources[source.source_id] = source
        collections[source.source_id] = Collection.model_validate(raw["collection"])
    if set(sources) != set(source_ids):
        problems.append("快照与报告来源集合不一致")
    if manifest["window"] not in {"all", "day", "week", "month"}:
        raise ValueError("未知窗口")
    window = cast(Window, manifest["window"])
    as_of = datetime.fromisoformat(manifest["as_of"])
    if as_of.tzinfo is None:
        raise ValueError("筛选时间缺少时区")
    undated = (
        manifest.get("configuration", {}).get("incremental", {}).get("include_undated_new", True)
    )
    companies: set[str] = set()
    excluded_companies: set[str] = set()
    raw_keys: set[str] = set()
    for row in reports:
        if row.source_id not in sources:
            continue
        collection = collections[row.source_id]
        for job in collection.jobs:
            if job.source_id != row.source_id or job.key in raw_keys:
                problems.append(f"岗位身份不符或重复: {job.key}")
            raw_keys.add(job.key)
        selected = sum(
            select(j, window, as_of, include_undated_new=undated).kind == "include"
            for j in collection.jobs
        )
        if row.collected != len(collection.jobs) or row.selected != selected:
            problems.append(f"原始/有效岗位计数不符: {row.source_id}")
        if row.coverage != "complete" or collection.coverage_issues() or row.issues:
            problems.append(f"目录覆盖未通过: {row.source_id}")
            excluded_companies.add(sources[row.source_id].company_id)
        elif selected > 0:
            companies.add(sources[row.source_id].company_id)
    companies -= excluded_companies
    if expected_companies is not None and len(companies) != expected_companies:
        problems.append(f"有岗位企业数应为 {expected_companies}，实际 {len(companies)}")
    delivered = [
        Job.model_validate(row) for row in json.loads((root / "delivered.json").read_text())
    ]
    delivered_keys = [job.key for job in delivered]
    if len(delivered_keys) != len(set(delivered_keys)) or not set(delivered_keys) <= raw_keys:
        problems.append("导出岗位身份重复或不属于原始快照")
    by_source = Counter(job.source_id for job in delivered)
    for row in reports:
        if by_source[row.source_id] != row.exported + row.review:
            problems.append(f"导出快照数量不符: {row.source_id}")
    counts = {
        name: sum(getattr(row, name) for row in reports)
        for name in ("collected", "selected", "exported", "review")
    }
    for name in ("collected", "exported", "review"):
        if report[name] != counts[name]:
            problems.append(f"报告汇总计数不符: {name}")
    workbook_rows: dict[str, int] = {}
    for filename, expected in [
        ("岗位信息汇总.xlsx", len(delivered)),
        ("jobs.xlsx", counts["exported"]),
    ]:
        book = load_workbook(root / filename, read_only=True, data_only=False)
        try:
            sheet = book.worksheets[0]
            if sheet.max_column != 13:
                problems.append(f"模板不是 13 列: {filename}")
            header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            if header[-2:] != ("截止时间", "发布时间"):
                problems.append(f"日期列标题不符: {filename}")
            count = 0
            for cells in sheet.iter_rows(min_row=2):
                if not any(cell.value is not None for cell in cells):
                    continue
                count += 1
                if any(cell.data_type == "f" for cell in cells):
                    problems.append(f"岗位数据包含公式: {filename}:{count + 1}")
                for cell in cells[11:13]:
                    value = cell.value
                    if value in (None, ""):
                        continue
                    try:
                        if (
                            not isinstance(value, str)
                            or len(value) != 10
                            or date.fromisoformat(value).isoformat() != value
                        ):
                            raise ValueError
                    except ValueError:
                        problems.append(f"日期格式错误: {filename}:{count + 1}")
            workbook_rows[filename] = count
            if count != expected:
                problems.append(f"表格行数不符: {filename} 应为 {expected}，实际 {count}")
        finally:
            book.close()
    return {
        "ok": not problems,
        "companies_with_jobs": len(companies),
        "company_ids": sorted(companies),
        "sources": len(reports),
        **counts,
        "information_rows": len(delivered),
        "workbook_rows": workbook_rows,
        "window": window,
        "as_of": manifest["as_of"],
        "errors": problems,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--expected-companies", type=int)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        if args.expected_companies is not None and args.expected_companies < 1:
            raise ValueError("目标企业数必须为正数")
        if args.output and args.output.resolve().is_relative_to(args.input.resolve()):
            raise ValueError("审计报告请保存到输入快照目录以外")
        result = audit(args.input, args.expected_companies)
        if args.output:
            if args.output.exists():
                raise ValueError("审计报告已存在，请使用新路径")
            write_json(args.output, result)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result["ok"] else 2
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print(f"无法核验交付: {type(exc).__name__}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
