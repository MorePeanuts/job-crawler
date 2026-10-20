"""Hunan Gold's 2027 Gold Cave recruitment-plan attachment."""

import hashlib
import io
import re
from typing import cast
from urllib.parse import urlsplit

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from job_crawler.adapters.base import Adapter, text
from job_crawler.adapters.platforms import check
from job_crawler.models import Collection, Scope
from job_crawler.storage import write_json
from job_crawler.transport import BrowserSession, CrawlError

EMPLOYER = "湖南黄金洞矿业有限责任公司"
ATTACHMENT = "湖南黄金洞矿业有限责任公司2027年招聘计划表.xlsx"
HEADER = ("序号", "岗位", "学历", "专业", "专业具体分支")
ANNOUNCEMENT_ID = "1012051"


def announcement(html: str) -> tuple[str, str]:
    soup = BeautifulSoup(html, "html.parser")
    page = text(soup.get_text(" ", strip=True))
    check(
        f"{EMPLOYER}招聘公告" in page
        and "招聘对象，寻找发光的你" in page
        and "应往届毕业生" in page
        and "招聘职位，精准匹配你的专业" in page,
        "黄金洞矿业公告主体或毕业生范围变化",
    )
    date = re.search(r"更新时间[：:]\s*(\d{4}-\d{2}-\d{2})", page)
    check(date is not None, "黄金洞矿业公告发布时间缺失")
    anchors = [a for a in soup.select("a[href]") if ATTACHMENT in text(a.get_text())]
    check(len(anchors) == 1, "黄金洞矿业招聘计划附件缺失或重复")
    url = str(anchors[0].get("href") or "")
    parts = urlsplit(url)
    check(
        parts.scheme == "https"
        and parts.hostname == "www.hnrcsc.com"
        and parts.path.endswith(".xlsx"),
        "黄金洞矿业招聘附件地址或文件格式变化",
    )
    assert date is not None
    return date[1], url


def positions(content: bytes) -> list[dict[str, str]]:
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        check(len(book.sheetnames) == 1, "黄金洞招聘计划工作表数量变化")
        rows = list(book.active.iter_rows(values_only=True)) if book.active is not None else []
        check(len(rows) >= 4, "黄金洞招聘计划无岗位")
        check(
            f"{EMPLOYER}2027年招聘计划表" in str(rows[1][0]).replace(" ", ""),
            "黄金洞招聘附件主体或届别变化",
        )
        check(tuple(rows[2][:5]) == HEADER, "黄金洞招聘计划列含义变化")
        result: list[dict[str, str]] = []
        for number, row in enumerate(rows[3:], 1):
            check(len(row) == 5 and row[0] == number, "黄金洞招聘计划编号不连续或列数变化")
            title, education, category, majors = (text(item) for item in row[1:5])
            check(all((title, education, category, majors)), "黄金洞招聘计划岗位字段缺失")
            result.append(
                {
                    "id": str(number),
                    "title": title,
                    "education": education,
                    "category": category,
                    "majors": majors,
                }
            )
        return result
    except (OSError, ValueError, TypeError) as exc:
        raise CrawlError("schema_changed", f"黄金洞招聘计划 XLSX 无法解析：{exc}") from exc


class HngoldGov2027(Adapter):
    name = "hngold-gov-2027"
    host = "qz.hnrcsc.com"
    path = "/ptemp/"
    scope_description = "湖南人才网黄金洞矿业 2027 年招聘计划附件全部岗位"

    async def collect(self) -> Collection:
        try:
            parts = urlsplit(self.source.url)
            check(
                parts.hostname == self.host
                and parts.path == self.path
                and f"id={ANNOUNCEMENT_ID}" in parts.fragment,
                "湖南人才网黄金洞公告入口变化",
            )
            await self.session.open(
                self.source.url,
                ready_endpoints=("/hnrcwzp/person-service/portalhomepage/getDetails",),
            )
            await self.harvest()
        except Exception as exc:
            self.error(exc)
        return self.result

    async def harvest(self) -> None:
        session = cast(BrowserSession, self.session)
        check(session.page is not None, "湖南人才网公告未就绪")
        published, attachment_url = announcement(await session.page.content())
        announcement_evidence = session.last_evidence
        session.consume_request()
        response = await session.page.context.request.get(attachment_url, timeout=30000)
        check(response.ok, f"黄金洞招聘计划附件 HTTP {response.status}")
        content = await response.body()
        check(0 < len(content) <= 2_000_000, "黄金洞招聘计划附件大小异常")
        rows = positions(content)
        evidence_path = session.output / f"attachment-{ANNOUNCEMENT_ID}.json"
        write_json(
            evidence_path,
            {
                "url": attachment_url,
                "sha256": hashlib.sha256(content).hexdigest(),
                "byte_count": len(content),
                "rows": rows,
            },
        )
        scope = Scope(name=f"hngold-2027-{ANNOUNCEMENT_ID}", expected=len(rows))
        self.result.scopes.append(scope)
        for row in self.add_page(scope, rows, set(), lambda item: item["id"]):
            job = self.job(
                row["id"],
                title=row["title"],
                source_url=self.source.url,
                unit=EMPLOYER,
                work_type="校招",
                education=row["education"],
                requirements=f"{row['category']}；{row['majors']}",
                published=published,
                field_sources={
                    "unit": "湖南人才网公告主体及上市代码 002155；仅黄金洞子公司",
                    "title": "公告唯一 2027 年招聘计划附件岗位列",
                    "education": "招聘计划学历列",
                    "requirements": "招聘计划专业及分支列",
                    "published": "湖南人才网公告更新时间；不是逐岗发布日期",
                    "location": "source_missing:公告公司地址不是逐岗工作地点",
                    "duties": "source_missing:公告未给逐岗职责",
                    "deadline": "source_missing:公告未给投递截止日期",
                },
            )
            job.evidence = [announcement_evidence, str(evidence_path)]
            self.result.jobs.append(job)
            scope.details += 1
        scope.terminal = True
