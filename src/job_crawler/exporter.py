"""Template-compatible output with a separate, lossless review channel."""

from copy import copy
from importlib.resources import files
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from job_crawler.config import MappingConfig
from job_crawler.models import Job

HEADERS = [
    "公司",
    "岗位名称",
    "工作类型",
    "一级分类",
    "二级分类",
    "学历",
    "地点",
    "详细地址",
    "岗位职责",
    "任职要求",
    "申请链接",
]
FIELDS = [
    "company",
    "title",
    "work_type",
    "category_primary",
    "category_secondary",
    "education",
    "location",
    "address",
    "duties",
    "requirements",
    "apply_url",
]


def template_path() -> Path:
    return Path(str(files("job_crawler").joinpath("assets/job-template.xlsx")))


def enums() -> dict[str, set[str]]:
    workbook = load_workbook(template_path(), read_only=True)
    sheet = workbook["枚举值"]
    result = {
        name: {
            str(row[0])
            for row in sheet.iter_rows(min_row=2, min_col=i, max_col=i, values_only=True)
            if row[0]
        }
        for i, name in enumerate(["work_type", "primary", "secondary", "education"], 1)
    }
    workbook.close()
    return result


def validate_mapping(mapping: MappingConfig) -> None:
    allowed = enums()
    for primary, children in mapping.category_paths.items():
        if primary not in allowed["primary"] or not set(children) <= allowed["secondary"]:
            raise ValueError(f"分类路径不属于模板枚举: {primary}")
    for raw, pair in mapping.categories.items():
        if len(pair) != 2 or pair[1] not in mapping.category_paths.get(pair[0], []):
            raise ValueError(f"分类映射未在分类树声明: {raw}")
    if not set(mapping.education.values()) <= allowed["education"]:
        raise ValueError("学历映射目标必须来自模板枚举")


def normalize(job: Job, mapping: MappingConfig) -> Job:
    result = job.model_copy(deep=True)
    pair = mapping.categories.get(job.category_raw)
    if pair:
        result.category_primary, result.category_secondary = pair
        result.field_sources["category"] = f"config:mapping.categories:{job.category_raw}"
    result.education = mapping.education.get(job.education, job.education)
    result.location = mapping.locations.get(job.location, job.location)
    if not result.apply_url and mapping.use_source_url_as_apply_url:
        result.apply_url = result.source_url
        result.field_sources["apply_url"] = "source_url:原始岗位详情入口，不保证免登录投递"
    return result


def import_errors(job: Job, mapping: MappingConfig, allowed: dict[str, set[str]]) -> list[str]:
    errors: list[str] = []
    for field in (
        "company",
        "title",
        "work_type",
        "category_primary",
        "category_secondary",
        "education",
    ):
        if not getattr(job, field).strip():
            errors.append(f"missing:{field}")
    if job.work_type not in allowed["work_type"]:
        errors.append("invalid:work_type")
    if job.education not in allowed["education"]:
        errors.append("invalid:education")
    if job.category_secondary not in mapping.category_paths.get(job.category_primary, []):
        errors.append("invalid:category_path")
    # Never invent the missing district; province/city or municipality/district is enough.
    parts = job.location.split("/")
    if (
        len(parts) not in {2, 3}
        or any(not p.strip() for p in parts)
        or any(c in job.location for c in "、,;；")
    ):
        errors.append("invalid:location_needs_mapping")
    if job.apply_url and urlsplit(job.apply_url).scheme not in {"http", "https"}:
        errors.append("invalid:apply_url")
    for name in FIELDS:
        value = getattr(job, name)
        if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
            errors.append(f"invalid:excel_text:{name}")
    return errors


def export_workbook(path: Path, jobs: list[Job], *, information: bool = False) -> None:
    workbook = load_workbook(template_path())
    sheet = workbook["岗位导入"]
    if [cell.value for cell in sheet[1]] != HEADERS:
        raise ValueError("模板表头与导出契约不一致")
    styles = [copy(cell._style) for cell in sheet[2]]
    sheet.delete_rows(2, sheet.max_row)
    if information:
        sheet.title = "岗位信息（待补全）"
        workbook["填写说明"].append(
            [
                "此表含待补全或待核验信息，不保证可直接导入。采集时刻和覆盖报告见 report.json；严格导入表为 jobs.xlsx。"
            ]
        )
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = f"A1:K{max(1, len(jobs) + 1)}"
    for row_index, job in enumerate(jobs, 2):
        for col_index, field in enumerate(FIELDS, 1):
            value = getattr(job, field)
            if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
                raise ValueError(f"无法无损导出 {job.key}:{field}")
            cell = sheet.cell(row_index, col_index, value)
            cell.data_type = "s"  # Untrusted website text must never become an Excel formula.
            cell._style = copy(styles[col_index - 1])
    for validation in sheet.data_validations.dataValidation:
        column = str(validation.sqref).split("2:")[0]
        validation.sqref = f"{column}2:{column}{max(1000, len(jobs) + 1)}"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
